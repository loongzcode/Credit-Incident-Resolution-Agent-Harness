"""Bounded XLSX inventory parsing. Never infer tools, authority or permissions."""
from io import BytesIO
from pathlib import PurePath
from zipfile import ZipFile
from hashlib import sha256
from openpyxl import load_workbook
from .models import CompanySystemInventory, SourceMetadata, AdminError

HEADERS = {
    "system_code": ("system_code", "系统编码", "系统英文名", "系统英文名称", "英文名称"),
    "system_name": ("system_name", "系统名称", "系统中文名", "系统中文名称", "中文名称"),
    "system_category": ("system_category", "系统分类", "系统类别"),
    "description": ("description", "系统简介", "简介", "系统描述"),
    "main_functions": ("main_functions", "主要功能", "主要功能说明"),
}


def parse_inventory(content, filename, actor, now):
    if not filename or not filename.lower().endswith(".xlsx") or len(content) > 5_000_000:
        raise AdminError("INVALID_WORKBOOK", 422)
    # Compressed payloads are bounded before the XML parser touches them.
    try:
        with ZipFile(BytesIO(content)) as z:
            if len(z.infolist()) > 300 or sum(i.file_size for i in z.infolist()) > 25_000_000:
                raise ValueError()
            if any("vbaproject" in i.filename.lower() for i in z.infolist()):
                raise ValueError()
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        rows, invalid, seen = [], [], set()
        try:
            for sheet in book:
                stream = sheet.iter_rows()
                first = next(stream, ())
                if len(first) > 100:
                    raise ValueError()
                headers = [str(c.value or "").strip() for c in first]
                indices = {k: next((i for i, h in enumerate(headers) if h in names), None) for k, names in HEADERS.items()}
                if indices["system_code"] is None or indices["system_name"] is None:
                    invalid.append({"row_ref": sheet.title + "!1", "code": "REQUIRED_HEADERS_MISSING"})
                    continue
                for n, cells in enumerate(stream, 2):
                    if n > 5001 or len(cells) > 100:
                        raise ValueError()
                    if all(c.value is None for c in cells):
                        continue
                    ref = f"{sheet.title}!{n}"
                    try:
                        if any(c.data_type == "f" for c in cells):
                            raise ValueError()
                        values = {k: str(cells[i].value or "").strip() if i is not None and i < len(cells) else ""
                                  for k, i in indices.items()}
                        row = CompanySystemInventory(**values, source_metadata=SourceMetadata(
                            source_file_name=PurePath(filename.replace('\\', '/')).name[:200], source_file_hash=sha256(content).hexdigest(),
                            source_row_ref=ref, imported_at=now, imported_by=actor))
                        if row.system_code in seen:
                            raise ValueError()
                        seen.add(row.system_code)
                        rows.append(row.model_dump(mode="json"))
                    except ValueError:
                        invalid.append({"row_ref": ref, "code": "INVALID_OR_DUPLICATE_ROW"})
        finally:
            book.close()
        if len(rows) > 5000 or len(invalid) > 5000:
            raise ValueError()
        return rows, invalid
    except AdminError:
        raise
    except Exception:
        raise AdminError("INVALID_WORKBOOK", 422) from None
