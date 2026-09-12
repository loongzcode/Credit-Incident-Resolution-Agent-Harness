"""Production-shaped workbooks with wholly synthetic descriptions and functions."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZipFile
import pytest
from openpyxl import Workbook
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from tests.test_registry_admin import console, routed, harness, request, act, approved, draft
from credit_harness.registry_admin.importer import parse_inventory
from credit_harness.registry_admin.models import AdminError, CompanySystemRecord
from credit_harness.registry_admin.inventory import read_records
from credit_harness.registry_admin.tables import InventoryHeadRow, InventoryVersionRow, AdminAuditRow
from credit_harness.registry.tables import RegistryAuditRow, RegistryVersionRow

GROUPS = ("九里云系统", "融担系统")


def synthetic_book(*, conflict=False, wrong_group=False, large=False, merged=True):
    book = Workbook(); book.remove(book.active)
    for index, group in enumerate(GROUPS):
        sheet = book.create_sheet(group)
        sheet.append(["序号", "系统分类", "系统简称", "类型", "系统名称", "系统简要",
                      "系统主要功能" if index == 0 else "系统功能", "authority", "WRITE", "adapter"])
        codes = ["aut", "sso"] + ([f"SYN-{index}-{n:02}" for n in range(31)] if large else [f"SYN-{index}"])
        for n, code in enumerate(codes, 1):
            sheet.append([n, (GROUPS[1-index] if wrong_group else group) if n == 1 else None, code,
                "测试基础类型", "Synthetic " + code, "Synthetic summary " + ("conflict" if conflict and index == 1 and code == "aut" else code),
                "Synthetic function " + code, "AUTHORITATIVE", "yes", "should-not-import"])
        if merged:
            sheet.merge_cells(start_row=2, start_column=2, end_row=len(codes)+1, end_column=2)
    stream = BytesIO(); book.save(stream); book.close(); return stream.getvalue()


def parse(**kwargs):
    return parse_inventory(synthetic_book(**kwargs), "synthetic-company.xlsx", "synthetic-editor", datetime.now(timezone.utc))


def preview(x, **kwargs):
    response = request(x, "POST", "/import/preview", files={"file": ("synthetic-company.xlsx", synthetic_book(**kwargs))})
    assert response.status_code == 200
    return response.json()


def import_draft(x, **kwargs):
    p = preview(x, **kwargs)
    response = request(x, "POST", "/import/confirm", json=dict(preview_id=p["preview_id"], title="Synthetic import", reason="synthetic only"))
    assert response.status_code == 200
    return response.json()


def activate_inventory(x):
    cr = approved(x, import_draft(x))
    result = act(x, cr, "activate", "activator")
    assert result.status_code == 200
    return result.json()


def test_real_schema_headers_are_supported():
    p = parse(); assert not p["invalid"] and not p["conflicts"]
    assert {r["system_code"] for r in p["rows"]} == {"aut", "sso", "SYN-0", "SYN-1"}


def test_two_sheet_inventory_supported():
    p = parse(); assert p["source_sheet_count"] == 2 and p["groups"] == sorted(GROUPS)


@pytest.mark.parametrize("merged", [True, False])
def test_inventory_group_forward_filled(merged):
    p = parse(merged=merged)
    assert next(r for r in p["rows"] if r["system_code"] == "SYN-0")["memberships"][0]["inventory_group"] == GROUPS[0]


def test_system_type_comes_from_type_column():
    assert all(r["system_type_label"] == "测试基础类型" and "system_category" not in r for r in parse()["rows"])


def test_description_from_system_summary_column():
    assert all(r["description"] == "Synthetic summary " + r["system_code"] for r in parse()["rows"])


def test_both_function_column_names_supported():
    assert all(r["main_functions"] == "Synthetic function " + r["system_code"] for r in parse()["rows"])


def test_shared_aut_like_system_is_not_invalid_duplicate():
    p = parse(); assert not p["invalid"] and not p["conflicts"] and p["shared_systems"] == ["aut", "sso"]


def test_same_system_multiple_groups_has_one_master():
    rows = [r for r in parse()["rows"] if r["system_code"] == "aut"]
    assert len(rows) == 1
    assert {m["inventory_group"] for m in rows[0]["memberships"]} == set(GROUPS)
    assert {m["source_metadata"]["source_row_ref"] for m in rows[0]["memberships"]} == {g + "!2" for g in GROUPS}


def test_conflicting_same_code_fails_preview(console):
    p = preview(console, conflict=True)
    assert p["conflicts"][0]["code"] == "CONFLICTING_SYSTEM_DEFINITION"
    assert request(console, "POST", "/import/confirm", json=dict(preview_id=p["preview_id"], title="x", reason="x")).status_code == 422


def test_inventory_group_conflict_is_invalid():
    p = parse(wrong_group=True)
    assert {e["code"] for e in p["invalid"]} == {"INVALID_INVENTORY_GROUP"}


def test_import_preview_reports_source_and_unique_counts(console):
    p = preview(console, large=True)
    assert (p["source_rows"], p["unique_systems"], len(p["shared_systems"])) == (66, 64, 2)
    assert not p["invalid"] and not p["conflicts"]


def test_inventory_activation_does_not_reactivate_registry_version(console):
    x = console; original = x.repo.current()[0]
    snapshot = x.catalog.project(x.case)[1]
    with Session(x.cases.engine) as s:
        audit_count = s.scalar(select(func.count()).select_from(RegistryAuditRow))
        version_count = s.scalar(select(func.count()).select_from(RegistryVersionRow))
    cr = activate_inventory(x)
    assert cr["proposed_registry_version"] == original == x.repo.current()[0]
    x.catalog.revalidate_snapshot(x.case, snapshot)
    with Session(x.cases.engine) as s:
        assert s.scalar(select(func.count()).select_from(RegistryAuditRow)) == audit_count
        assert s.scalar(select(func.count()).select_from(RegistryVersionRow)) == version_count
        assert s.get(InventoryHeadRow, "demo").revision == 1
    event = request(x, "GET", "/audit", user="auditor").json()["items"][0]
    assert event["event"] == "INVENTORY_ACTIVATED" and event["after_inventory_revision"] == 1


@pytest.mark.parametrize("field", ["capabilities", "authority_rules", "systems"])
def test_inventory_never_creates_capability_authority_write_or_source(console, field):
    x = console; before = x.repo.current()[1].model_dump()[field]
    activate_inventory(x)
    assert x.repo.current()[1].model_dump()[field] == before


def bind(x, source_count=1):
    cr = draft(x, edited=False)
    definition = request(x, "GET", f'/change-requests/{cr["change_request_id"]}').json()["definition"]
    for source in definition["systems"][:source_count]:
        source["company_system_code"] = "aut"
    response = request(x, "PATCH", f'/change-requests/{cr["change_request_id"]}', json=dict(expected_revision=cr["revision"],
        title=cr["title"], reason=cr["reason"], definition=definition))
    assert response.status_code == 200
    cr = approved(x, response.json()); assert act(x, cr, "activate", "activator").status_code == 200


def test_company_system_detail_joins_agent_config(console):
    x = console; activate_inventory(x); bind(x)
    detail = request(x, "GET", "/systems/aut").json()
    assert detail["company_system"]["description"] == "Synthetic summary aut"
    assert detail["registry_sources"] and detail["capabilities"] and detail["authority"]
    assert "vault://" not in str(detail)


def test_unconfigured_company_system_visible(console):
    x = console; activate_inventory(x)
    result = request(x, "GET", "/systems").json()
    assert result["total"] == 4
    assert all(r["agent_configuration_status"] == "NOT_CONFIGURED" and r["capability_count"] == 0 for r in result["items"])


def test_one_company_system_can_bind_multiple_registry_sources(console):
    x = console; activate_inventory(x); bind(x, 2)
    detail = request(x, "GET", "/systems/aut").json()
    assert detail["agent_source_count"] == 2 and len(detail["registry_sources"]) == 2
    assert {s["company_system_code"] for s in detail["registry_sources"]} == {"aut"}


def test_binding_unknown_master_is_rejected(console):
    x = console; cr = draft(x, edited=False)
    d = request(x, "GET", f'/change-requests/{cr["change_request_id"]}').json()["definition"]
    d["systems"][0]["company_system_code"] = "MISSING"
    response = request(x, "PATCH", f'/change-requests/{cr["change_request_id"]}', json=dict(expected_revision=cr["revision"], title="x", reason="x", definition=d))
    assert response.json()["detail"] == "UNKNOWN_COMPANY_SYSTEM"


def test_company_list_filters_and_pagination(console):
    x = console; activate_inventory(x); bind(x)
    assert request(x, "GET", "/systems?system_code=aut&system_name=Synthetic&system_type_label=测试基础&inventory_group=融担系统&has_capability=true&configuration_status=CONFIGURED").json()["total"] == 1
    assert request(x, "GET", "/systems?has_capability=false").json()["total"] == 3
    assert len(request(x, "GET", "/systems?size=1&page=2").json()["items"]) == 1


def test_reject_requires_reason(console):
    x = console; cr = act(x, draft(x), "submit", "editor").json()
    for comment in (None, "", "  "):
        r = request(x, "POST", f'/change-requests/{cr["change_request_id"]}/reject', user="approver",
            json=dict(expected_revision=cr["revision"], decision_comment=comment))
        assert r.status_code == 422 and r.json()["detail"] == "REJECT_REASON_REQUIRED"


@pytest.mark.parametrize("action,field", [("approve", "approval_comment"), ("reject", "rejection_reason")])
def test_approval_comment_persisted(console, action, field):
    x = console; cr = act(x, draft(x), "submit", "editor").json()
    r = request(x, "POST", f'/change-requests/{cr["change_request_id"]}/{action}', user="approver",
        json=dict(expected_revision=cr["revision"], decision_comment="Synthetic plain decision"))
    assert r.status_code == 200 and r.json()[field] == "Synthetic plain decision"
    assert request(x, "GET", f'/change-requests/{cr["change_request_id"]}').json()["change_request"][field] == "Synthetic plain decision"
    assert request(x, "GET", "/audit", user="auditor").json()["items"][0]["decision_comment"] == "Synthetic plain decision"


@pytest.mark.parametrize("comment", ["<b>approve</b>", "x"*2001, "\x00invalid"])
def test_approval_rejects_html_or_unbounded_comment(console, comment):
    cr = act(console, draft(console), "submit", "editor").json()
    r = request(console, "POST", f'/change-requests/{cr["change_request_id"]}/approve', user="approver",
        json=dict(expected_revision=cr["revision"], decision_comment=comment))
    assert r.status_code == 422


def test_impact_distinguishes_changed_capability_types(console):
    x = console; cr = draft(x)  # description only; still globally fences snapshots
    path = f'/change-requests/{cr["change_request_id"]}'
    impact = request(x, "GET", path + "/impact").json()
    assert impact["changed_capability_types"] == []
    assert impact["runtime_revalidation_scope"] == "ALL_ACTIVE_REGISTERED_CASES"
    assert "affected_capability_types" not in impact
    d = request(x, "GET", path).json()["definition"]
    changed = d["capabilities"][0]; changed["adapter_id"] = "synthetic-changed-adapter"
    r = request(x, "PATCH", path, json=dict(expected_revision=cr["revision"], title="x", reason="x", definition=d))
    assert r.status_code == 200
    assert request(x, "GET", path + "/impact").json()["changed_capability_types"] == [changed["capability_type"]]


def test_inventory_impact_has_no_runtime_revalidation(console):
    cr = import_draft(console)
    impact = request(console, "GET", f'/change-requests/{cr["change_request_id"]}/impact').json()
    assert impact["runtime_revalidation_scope"] == "NONE" and impact["changed_capability_types"] == []


def test_inventory_concurrent_activation_only_one_wins(console):
    x = console; a, b = approved(x, import_draft(x)), approved(x, import_draft(x))
    with ThreadPoolExecutor(2) as pool:
        assert sorted(pool.map(lambda cr: act(x, cr, "activate", "activator").status_code, (a, b))) == [200, 409]
    with Session(x.cases.engine) as s:
        assert s.get(InventoryHeadRow, "demo").revision == 1
        assert s.scalar(select(func.count()).select_from(InventoryVersionRow)) == 1


def test_inventory_activation_audit_failure_rolls_back(console, monkeypatch):
    x = console; cr = approved(x, import_draft(x)); original = x.service._audit
    def fail(s, event, *args, **kwargs):
        original(s, event, *args, **kwargs)
        if event == "INVENTORY_ACTIVATED": raise RuntimeError("synthetic fault")
    monkeypatch.setattr(x.service, "_audit", fail)
    with pytest.raises(RuntimeError): act(x, cr, "activate", "activator")
    with Session(x.cases.engine) as s:
        assert s.get(InventoryHeadRow, "demo").revision == 0
        assert s.scalar(select(func.count()).select_from(InventoryVersionRow)) == 0


def test_legacy_inventory_read_migration_preserves_original():
    source = parse()["rows"][0]["memberships"][0]["source_metadata"]
    old = [dict(system_code="legacy", system_name="Synthetic legacy", system_category="Synthetic group",
                description="synthetic", main_functions="synthetic", source_metadata=source)]
    before = deepcopy(old); migrated = read_records(old)
    assert old == before and migrated[0]["system_type_label"] == ""
    assert migrated[0]["memberships"][0]["inventory_group"] == "Synthetic group"


@pytest.mark.parametrize("entry", ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"])
def test_workbook_active_content_rejected(entry):
    stream = BytesIO(synthetic_book())
    with ZipFile(stream, "a") as z: z.writestr(entry, "synthetic")
    with pytest.raises(AdminError, match="INVALID_WORKBOOK"):
        parse_inventory(stream.getvalue(), "synthetic.xlsx", "editor", datetime.now(timezone.utc))


def test_workbook_external_hyperlink_rejected():
    book = Workbook(); sheet = book.active
    sheet.append(["系统简称", "系统名称"]); sheet.append(["synthetic", "Synthetic"])
    sheet["B2"].hyperlink = "https://example.invalid/synthetic"
    stream = BytesIO(); book.save(stream); book.close()
    with pytest.raises(AdminError, match="INVALID_WORKBOOK"):
        parse_inventory(stream.getvalue(), "synthetic.xlsx", "editor", datetime.now(timezone.utc))


def test_membership_must_bind_same_company_system():
    row = parse()["rows"][0]; row["memberships"][0]["system_code"] = "OTHER"
    with pytest.raises(ValueError, match="membership must bind"):
        CompanySystemRecord.model_validate(row)


def test_reimport_same_content_is_unchanged_despite_new_source_time(console):
    x = console; activate_inventory(x); p = preview(x)
    assert len(p["unchanged"]) == 4 and not p["changed"] and not p["added"]


def test_partial_import_preserves_other_group_membership():
    from credit_harness.registry_admin.inventory import merge_records
    row = next(r for r in parse()["rows"] if r["system_code"] == "aut")
    partial = {**row, "memberships": row["memberships"][:1]}
    result = merge_records([row], [partial])[0]
    assert {m["inventory_group"] for m in result["memberships"]} == set(GROUPS)
