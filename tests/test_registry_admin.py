from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone, timedelta
from io import BytesIO
import json
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from tests.test_route_revisions import routed, harness
from credit_harness.registry_admin.identity import (EnterpriseIdentity, LocalIdentityProvider, ApplicationAccess, Role,
    OIDCIdentityProvider, IdentityError)
from credit_harness.registry_admin.service import RegistryAdministration, masked
from credit_harness.registry_admin.api import create_admin_app
from credit_harness.registry_admin.models import CreateDraft, EditDraft, AdminError, RollbackRequest
from credit_harness.registry_admin.tables import AdminAuditRow, InventoryHeadRow
from credit_harness.registry.tables import RegistryVersionRow
from credit_harness.cases.preconditions import AgentPreconditionFailed
from credit_harness.agent.revalidation import execution_precondition
from credit_harness.context.assembler import ReasoningContextAssembler
from tests.test_case_evidence import Q
from credit_harness.domain.enums import ToolName


@pytest.fixture
def console(routed):
    x = routed()
    now = datetime.now(timezone.utc)
    identities = {name: EnterpriseIdentity(user_id=name, display_name=name, groups=(name,), authenticated_at=now)
                  for name in ("viewer", "editor", "approver", "approver2", "activator", "auditor", "all")}
    mapping = {name: (role,) for name, role in zip(("viewer", "editor", "approver", "activator", "auditor"), Role)}
    mapping.update(approver2=(Role.REGISTRY_APPROVER,), all=tuple(Role))
    x.service = RegistryAdministration(x.cases.engine, "demo", ApplicationAccess(mapping))
    x.identities = identities
    provider = LocalIdentityProvider(identities, expires_at=now + timedelta(hours=2))
    x.api = TestClient(create_admin_app(x.service, provider, local_mode=True, allowed_origins=("https://admin.test",)))
    return x


def request(x, method, path="", *, user="editor", **kwargs):
    return x.api.request(method, "/admin/registry" + path,
        headers={"Authorization": "Bearer " + user, "X-Registry-Request": "1"}, **kwargs)


def draft(x, user="editor", edited=True):
    response = request(x, "POST", "/change-requests", user=user,
        json=dict(title="调整接入说明", reason="synthetic administration test", base_registry_version=x.repo.current()[0]))
    assert response.status_code == 200, response.text
    cr = response.json()
    if edited:
        definition = request(x, "GET", "/change-requests/" + cr["change_request_id"], user=user).json()["definition"]
        definition["systems"][0]["display_name"] += " updated"
        r = request(x, "PATCH", "/change-requests/" + cr["change_request_id"], user=user,
            json=dict(expected_revision=cr["revision"], title=cr["title"], reason=cr["reason"], definition=definition))
        assert r.status_code == 200, r.text
        cr = r.json()
    return cr


def act(x, cr, action, user):
    return request(x, "POST", f'/change-requests/{cr["change_request_id"]}/{action}', user=user,
                   json=dict(expected_revision=cr["revision"]))


def approved(x, cr=None):
    cr = cr or draft(x)
    r = act(x, cr, "submit", "editor"); assert r.status_code == 200, r.text
    r = act(x, r.json(), "approve", "approver"); assert r.status_code == 200, r.text
    return r.json()


def test_anonymous_cannot_access_admin(console):
    assert console.api.get("/admin/registry").status_code == 401


@pytest.mark.parametrize("user,action", [("viewer", "submit"), ("editor", "approve"), ("approver", "activate")])
def test_backend_role_permissions(console, user, action):
    cr = draft(console)
    assert act(console, cr, action, user).status_code == 403
    assert request(console, "POST", "/change-requests", user="viewer", json=dict(title="x", reason="x", base_registry_version=None)).status_code == 403


def test_sso_groups_map_to_roles(console):
    data = request(console, "GET", user="editor").json()
    assert data["roles"] == ["REGISTRY_EDITOR"]
    assert "REGISTRY_EDIT" in data["permissions"] and "REGISTRY_APPROVE" not in data["permissions"]


def test_creator_cannot_self_approve(console):
    cr = draft(console, user="all")
    cr = act(console, cr, "submit", "all").json()
    assert act(console, cr, "approve", "all").json()["detail"] == "SELF_APPROVAL_REJECTED"


@pytest.mark.parametrize("state", ["draft", "rejected"])
def test_unapproved_cannot_activate(console, state):
    cr = draft(console)
    if state == "rejected":
        cr = act(console, cr, "submit", "editor").json()
        cr = act(console, cr, "reject", "approver").json()
    assert act(console, cr, "activate", "activator").status_code == 409


def test_only_approved_activates_and_runtime_sees_only_active(console):
    x = console; before = x.repo.current()[0]; cr = approved(x)
    assert x.repo.current()[0] == before
    assert act(x, cr, "activate", "activator").status_code == 200
    assert x.repo.current()[0] == cr["proposed_registry_version"]


def test_active_version_immutable(console):
    x = console; before = x.repo.current()[0]
    with Session(x.cases.engine) as s:
        payload = s.get(RegistryVersionRow, before).payload
    cr = approved(x); act(x, cr, "activate", "activator")
    with Session(x.cases.engine) as s:
        assert s.get(RegistryVersionRow, before).payload == payload


def test_stale_base_version_rejected(console):
    x = console; a, b = approved(x), draft(x)
    assert act(x, a, "activate", "activator").status_code == 200
    assert act(x, b, "submit", "editor").json()["detail"] == "STALE_CHANGE_REQUEST"


def test_concurrent_activation_only_one_wins(console):
    x = console; a, b = approved(x), approved(x)
    def activate(cr):
        return act(x, cr, "activate", "activator").status_code
    with ThreadPoolExecutor(2) as pool:
        assert sorted(pool.map(activate, (a, b))) == [200, 409]


def test_concurrent_approvals_state_cas(console):
    x = console; cr = act(x, draft(x), "submit", "editor").json()
    with ThreadPoolExecutor(2) as pool:
        results = list(pool.map(lambda actor: act(x, cr, "approve", actor).status_code, ("approver", "approver2")))
    assert sorted(results) == [200, 409]


def test_activation_and_audit_are_atomic(console, monkeypatch):
    x = console; before = x.repo.current()[0]; cr = approved(x)
    original = x.service._audit
    def fail(s, event, *args, **kwargs):
        original(s, event, *args, **kwargs)
        if event == "ACTIVATED":
            raise RuntimeError("synthetic commit failure")
    monkeypatch.setattr(x.service, "_audit", fail)
    with pytest.raises(RuntimeError):
        act(x, cr, "activate", "activator")
    assert x.repo.current()[0] == before
    assert x.service.get(x.identities["viewer"], cr["change_request_id"])["change_request"]["status"] == "APPROVED"


def test_rollback_creates_audited_change_request(console):
    x = console; original = x.repo.current()[0]; cr = approved(x)
    act(x, cr, "activate", "activator")
    r = request(x, "POST", "/rollback", json=dict(title="回退", reason="恢复已审核配置", base_registry_version=x.repo.current()[0], target_version=original))
    assert r.status_code == 200, r.text
    rollback = approved(x, r.json()); assert rollback["kind"] == "ROLLBACK"
    assert act(x, rollback, "activate", "activator").status_code == 200
    assert x.repo.current()[0] == original
    events = [e["event"] for e in request(x, "GET", "/audit?size=100", user="auditor").json()["items"]]
    assert "ROLLBACK_REQUESTED" in events and "ROLLED_BACK" in events


def test_version_diff_deterministic_and_payment_high_risk(console):
    x = console; cr = draft(x)
    a = request(x, "GET", "/change-requests/" + cr["change_request_id"]).json()
    definition = a["definition"]
    payment = next(c for c in definition["capabilities"] if c["tool_name"] == ToolName.PAYMENT.value)
    payment["adapter_id"] = "synthetic-new-payment"
    r = request(x, "PATCH", "/change-requests/" + cr["change_request_id"], json=dict(expected_revision=cr["revision"],
        title=cr["title"], reason=cr["reason"], definition=definition))
    assert r.status_code == 200, r.text
    cr = r.json(); path = f'/versions/{cr["base_registry_version"]}/diff/{cr["proposed_registry_version"]}'
    diff = request(x, "GET", path).json()
    assert diff == request(x, "GET", path).json()
    assert diff["risk_level"] == "HIGH_RISK" and "PAYMENT_FINALITY_SOURCE_CHANGE" in diff["risk_reasons"]


def test_impact_analysis_read_only_and_paginated(console):
    x = console; before_case = x.cases.get(x.case.case_id); head = x.repo.current()[0]; cr = draft(x)
    snapshot = x.catalog.project(before_case)[1]
    r = request(x, "GET", f'/change-requests/{cr["change_request_id"]}/impact?size=1').json()
    assert r["affected_active_case_count"] == 1 and len(r["case_page"]["items"]) == 1
    assert r["invalidated_snapshot_count"] == 1
    assert x.cases.get(x.case.case_id) == before_case and x.repo.current()[0] == head
    x.catalog.revalidate_snapshot(before_case, snapshot)


def test_impact_includes_removed_capability_type(console):
    x = console; cr = draft(x)
    definition = request(x, "GET", "/change-requests/" + cr["change_request_id"]).json()["definition"]
    removed = {c["capability_id"] for c in definition["capabilities"] if c["tool_name"] == ToolName.ACCOUNTING.value}
    definition["capabilities"] = [c for c in definition["capabilities"] if c["capability_id"] not in removed]
    definition["authority_rules"] = [r for r in definition["authority_rules"] if r["capability_id"] not in removed]
    r = request(x, "PATCH", "/change-requests/" + cr["change_request_id"], json=dict(expected_revision=cr["revision"],
        title=cr["title"], reason=cr["reason"], definition=definition))
    assert r.status_code == 200, r.text
    impact = request(x, "GET", f'/change-requests/{cr["change_request_id"]}/impact').json()
    assert "READ_ACCOUNTING_ENTRY" in impact["affected_capability_types"]


def workbook(extra=None):
    from openpyxl import Workbook
    book = Workbook(); sheet = book.active
    sheet.append(["系统编码", "系统中文名", "系统分类", "简介", "主要功能", "authority", "WRITE"])
    sheet.append(["PAY-DEMO", "测试公司支付清单", "支付", "synthetic", "查询", "AUTHORITATIVE", "yes"])
    if extra:
        sheet.append(extra)
    output = BytesIO(); book.save(output); return output.getvalue()


def import_preview(x, content=None):
    return request(x, "POST", "/import/preview", files={"file": ("系统中英文名对照及简介-v20240715bylibo.xlsx", content or workbook(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


def test_excel_preview_does_not_mutate_registry(console):
    x = console; before = x.repo.current(); data = import_preview(x).json()
    assert data["added"] == ["PAY-DEMO"] and not data["invalid"]
    assert x.repo.current() == before
    assert request(x, "GET", "/inventory").json()["total"] == 0


def test_excel_import_preserves_capabilities_and_cannot_create_authority_or_write(console):
    x = console; before = x.repo.current(); p = import_preview(x).json()
    r = request(x, "POST", "/import/confirm", json=dict(preview_id=p["preview_id"], title="导入清单", reason="仅公司资料"))
    assert r.status_code == 200, r.text
    cr = approved(x, r.json()); act(x, cr, "activate", "activator")
    assert x.repo.current() == before
    row = request(x, "GET", "/inventory").json()["items"][0]
    assert "authority" not in row and "WRITE" not in row
    assert row["source_metadata"]["source_row_ref"].endswith("!2")
    assert len(row["source_metadata"]["source_file_hash"]) == 64


def test_excel_invalid_rows_cannot_confirm(console):
    x = console; p = import_preview(x, workbook(["PAY-DEMO", "duplicate"])).json()
    assert p["invalid"]
    assert request(x, "POST", "/import/confirm", json=dict(preview_id=p["preview_id"], title="x", reason="x")).status_code == 422


def test_excel_formula_rejected(console):
    p = import_preview(console, workbook(["EVIL", "=HYPERLINK(\"https://example.com\")"])).json()
    assert p["invalid"]


def test_secrets_absent_from_api_and_credential_refs_masked(console):
    x = console; cr = draft(x)
    data = request(x, "GET", "/change-requests/" + cr["change_request_id"]).json()
    assert "vault://" not in json.dumps(data)
    assert any(s["credential_ref"] == "[MASKED]" for s in data["definition"]["systems"])
    data["definition"]["systems"][0]["credential_ref"] = "secret-value-never-echo"
    r = request(x, "PATCH", "/change-requests/" + cr["change_request_id"], json=dict(expected_revision=cr["revision"],
        title="x", reason="x", definition=data["definition"]))
    assert r.status_code == 422 and "secret-value" not in r.text
    assert "vault://" not in request(x, "GET", "/audit", user="auditor").text


def test_audit_append_only_api(console):
    x = console; draft(x)
    items = request(x, "GET", "/audit", user="auditor").json()["items"]
    assert items and all(e["actor"] and e["request_id"] and e["time"] for e in items)
    assert request(x, "PATCH", "/audit", user="all", json={}).status_code == 405
    assert request(x, "DELETE", "/audit", user="all").status_code == 405


def test_old_capability_snapshot_stales_after_admin_activation(console):
    x = console; case = x.cases.get(x.case.case_id)
    snapshot = ReasoningContextAssembler(catalog=x.catalog).build(case, ())
    cr = approved(x); act(x, cr, "activate", "activator")
    with pytest.raises(AgentPreconditionFailed, match="STALE_CAPABILITY"):
        x.executor.execute_if_current(case.case_id, ToolName.PAYMENT, Q, precondition=execution_precondition(case, snapshot))


def test_csrf_and_role_spoofing_rejected(console):
    x = console; body = dict(title="x", reason="x", base_registry_version=x.repo.current()[0])
    assert x.api.post("/admin/registry/change-requests", headers={"Authorization": "Bearer editor"}, json=body).status_code == 403
    assert x.api.post("/admin/registry/change-requests", headers={"Authorization": "Bearer editor", "X-Registry-Request": "1", "Origin": "https://evil.test"}, json=body).status_code == 403
    assert x.api.get("/admin/registry", headers={"X-Role": "REGISTRY_ACTIVATOR"}).status_code == 401


def test_local_identity_cannot_be_production_default(console):
    with pytest.raises(ValueError):
        create_admin_app(console.service, LocalIdentityProvider({}, expires_at=datetime.now(timezone.utc)))


def test_rollback_rejects_never_activated_candidate(console):
    x = console; cr = draft(x)
    r = request(x, "POST", "/rollback", json=dict(title="invalid rollback", reason="test", base_registry_version=x.repo.current()[0],
        target_version=cr["proposed_registry_version"]))
    assert r.status_code == 422 and r.json()["detail"] == "ROLLBACK_TARGET_NEVER_ACTIVATED"
    versions = request(x, "GET", "/versions").json()["items"]
    assert versions[0]["version"] == cr["proposed_registry_version"]
    assert versions[0]["registered_at"] and not versions[0]["activated_before"]


def test_stale_submitted_cannot_be_approved(console):
    x = console; cr = act(x, draft(x), "submit", "editor").json(); other = approved(x)
    act(x, other, "activate", "activator")
    assert act(x, cr, "approve", "approver").json()["detail"] == "STALE_CHANGE_REQUEST"


def test_foreign_tenant_admin_cannot_read_or_edit_cr(console):
    x = console; cr = draft(x)
    foreign = RegistryAdministration(x.cases.engine, "foreign-admin", x.service.access)
    with pytest.raises(AdminError, match="CHANGE_REQUEST_NOT_FOUND"):
        foreign.get(x.identities["viewer"], cr["change_request_id"])
    with pytest.raises(AdminError, match="VERSION_NOT_FOUND"):
        foreign.read(x.identities["viewer"], "definition", version=cr["proposed_registry_version"])


def test_inventory_preview_review_and_idempotent_confirm(console):
    x = console; p = import_preview(x).json(); body = dict(preview_id=p["preview_id"], title="inventory", reason="test")
    a = request(x, "POST", "/import/confirm", json=body).json()
    b = request(x, "POST", "/import/confirm", json=body).json()
    assert a == b
    details = request(x, "GET", "/change-requests/" + a["change_request_id"]).json()
    assert details["inventory_changes"][0]["system_code"] == "PAY-DEMO"
    assert request(x, "GET", "/inventory").json()["total"] == 0


@pytest.mark.parametrize("action,expected", [("cancel", "CANCELED"), ("supersede", "SUPERSEDED")])
def test_terminal_change_cannot_reenter_edit_or_activation(console, action, expected):
    x = console; cr = act(x, draft(x), action, "editor").json()
    assert cr["status"] == expected
    assert act(x, cr, "submit", "editor").status_code == 409
    assert act(x, cr, "activate", "activator").status_code == 409


def test_diff_never_exposes_changed_credential_reference(console):
    x = console
    from credit_harness.registry_admin.diff import version_diff
    before = x.repo.current()[1]
    after = before.model_copy(update=dict(systems=(before.systems[0].model_copy(update=dict(credential_ref="vault://hidden/new")), *before.systems[1:])))
    serialized = version_diff(before, after).model_dump_json()
    assert "vault://" not in serialized and "[MASKED]" in serialized


def test_excel_upload_invalid_zip_fails_closed(console):
    assert import_preview(console, b'not a zip').status_code == 422


def test_expired_local_token_rejected():
    now = datetime.now(timezone.utc)
    provider = LocalIdentityProvider({'x': EnterpriseIdentity(user_id='x', display_name='x', groups=(), authenticated_at=now)}, expires_at=now)
    with pytest.raises(IdentityError): provider.authenticate('x')


@pytest.mark.parametrize("invalid", [None, "issuer", "audience", "expired", "scope", "signature"])
def test_oidc_offline_signature_claim_validation(invalid):
    import jwt
    from cryptography.hazmat.primitives.asymmetric import rsa
    from types import SimpleNamespace
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    now = datetime.now(timezone.utc)
    provider = OIDCIdentityProvider(issuer="https://sso.test", audience="registry-api", jwks_url="https://sso.test/jwks")
    provider.keys = SimpleNamespace(get_signing_key_from_jwt=lambda _: SimpleNamespace(key=key.public_key()))
    claims = dict(iss="https://sso.test", aud="registry-api", sub="user-1", iat=now, exp=now+timedelta(minutes=2),
                  scope="registry.admin", groups=["registry-editors"])
    if invalid == "issuer": claims["iss"] = "https://evil.test"
    if invalid == "audience": claims["aud"] = "investigation-api"
    if invalid == "expired": claims["exp"] = now-timedelta(minutes=2)
    if invalid == "scope": claims["scope"] = "unrelated"
    signing = rsa.generate_private_key(public_exponent=65537, key_size=2048) if invalid == "signature" else key
    token = jwt.encode(claims, signing, algorithm="RS256", headers={"kid": "test"})
    if invalid:
        with pytest.raises(IdentityError): provider.authenticate(token)
    else:
        assert provider.authenticate(token).groups == ("registry-editors",)
